#!/usr/bin/env python3
"""Build `supplementary_tables.xlsx`: every supplementary and appendix table of the manuscript, one sheet each.

`manuscript.md` names, in the paragraphs headed "Supplementary Table S1" to "S9" and "Appendix Table A1" to "A9", the
CSV and TSV files of `data/` that make up each table. Those files are named by what they hold
(`criteria_firing_v37.csv`) and not by the number of the table, which makes the tables hard to find in a repository,
so this script gathers them into one workbook: a sheet "Contents", then a sheet for each table, with the files of that
table one under another, each under its own file name and with the values as the file holds them. A number is written
to 16 significant digits, which is all that openpyxl writes and one more than a spreadsheet shows, so a value that
needs 17 can differ from its file in the last digit; the check allows a relative difference of 1e-15 and no more.

The assignment of files to tables is read from the manuscript and not repeated here, so the workbook cannot disagree
with it, and the build stops if a file is in no table, is in two, or does not exist.

A cell is typed from the text of the file and not by pandas, so that "NA", "None" or "007" stays what it was written
as and an integer column with a blank in it stays integer.

ClinVar's 2-star+ classifications are a reference and not ground truth.

Usage
  build_supplement.py --write    write the workbook, unless it already holds what data/ and manuscript.md now say
  build_supplement.py --check    exit non-zero if the workbook differs from what data/ and manuscript.md now say
"""

from __future__ import annotations

import argparse
import csv
import math
import re
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink

HERE = Path(__file__).resolve().parent
DATA = HERE / "data"
MANUSCRIPT = HERE / "manuscript.md"
WORKBOOK = HERE / "supplementary_tables.xlsx"

LABEL = re.compile(r"\*\*((?:Supplementary|Appendix) Table ([SA]\d+))\.\*\*[ \t]*")
FILES = re.compile(r"`([\w.\-]+\.(?:csv|tsv))`")
INT = re.compile(r"-?(?:0|[1-9]\d*)")
FLOAT = re.compile(r"-?(?:0|[1-9]\d*)?\.\d+(?:[eE][+-]?\d+)?|-?(?:0|[1-9]\d*)(?:\.\d+)?[eE][+-]?\d+")

INTRO = (
    "The tables of the manuscript on building, reviewing and measuring the fastVEP ACMG/AMP classifier. "
    "Supplementary Tables S1 to S9 are the nine tables that summarise the data best, and Appendix Tables A1 to A9 hold "
    "every other table. Each sheet holds the CSV or TSV files of `manuscript/acmg/data/` that make up its table, one "
    "under another, each under its file name.\n"
    "ClinVar 2-star+: records of review status two stars or better (criteria provided, multiple submitters, no "
    "conflicts; reviewed by expert panel; practice guideline), release of 4 April 2026 unless a table says October 2026. "
    "P/LP: Pathogenic or Likely pathogenic; B/LB: Likely benign or Benign; VUS: uncertain significance. "
    "ClinVar's classifications are a reference and not ground truth."
)

MAX_WIDTH = 60  # characters; a longer value is not wrapped, it runs on
CHARS_PER_UNIT = 1.2  # characters of Calibri 11 that fit in one unit of column width, a little fewer than they do in practice
CAPTION_COLUMNS = 8  # the caption of a table is merged across this many columns, or all of them if there are fewer


def tables() -> list[dict]:
    """The supplementary and appendix tables of the manuscript, in its order: name, id, caption and files."""
    back = MANUSCRIPT.read_text().split("\n## Figure legends\n", 1)[1]
    found = []
    for para in back.split("\n\n"):
        m = LABEL.match(para)
        if m:
            found.append(
                {
                    "name": m.group(1),
                    "id": m.group(2),
                    "caption": para[m.end() :].replace("`", ""),
                    "files": list(dict.fromkeys(FILES.findall(para))),
                }
            )
    listed = [f for t in found for f in t["files"]]
    present = sorted(p.name for p in DATA.iterdir() if p.suffix in (".csv", ".tsv"))
    twice = sorted({f for f in listed if listed.count(f) > 1})
    problems = []
    if twice:
        problems.append(f"in two tables: {twice}")
    if sorted(set(listed) - set(present)):
        problems.append(f"named in the manuscript and not found in data/: {sorted(set(listed) - set(present))}")
    if sorted(set(present) - set(listed)):
        problems.append(f"in data/ and in no table: {sorted(set(present) - set(listed))}")
    if not found or problems:
        sys.exit("manuscript.md and data/ disagree about the tables: " + "; ".join(problems or ["no table found"]))
    return found


def typed(text: str):
    """A cell of a data file as a spreadsheet value: blank, boolean, integer, float or the text itself."""
    if text == "":
        return None
    if text in ("True", "False"):
        return text == "True"
    if INT.fullmatch(text):
        return int(text)
    if FLOAT.fullmatch(text):
        return float(text)
    return text


def read_file(name: str) -> tuple[list[str], list[list]]:
    with open(DATA / name, newline="", encoding="utf-8") as fh:
        rows = list(csv.reader(fh, delimiter="\t" if name.endswith(".tsv") else ","))
    header, body = rows[0], rows[1:]
    ragged = [i + 2 for i, r in enumerate(body) if len(r) != len(header)]
    if ragged:
        sys.exit(f"{name}: rows {ragged[:5]} do not have {len(header)} fields")
    return header, [[typed(c) for c in r] for r in body]


# --------------------------------------------------------------------------------- what each sheet holds


def table_rows(t: dict) -> list[tuple[str, list]]:
    """A table's sheet as (kind, values) rows; the kind decides the style and is not written."""
    rows: list[tuple[str, list]] = [("title", [t["name"]]), ("caption", [t["caption"]]), ("blank", [])]
    for name in t["files"]:
        header, body = read_file(name)
        rows.append(("file", [f"{name} ({len(body):,} rows, {len(header)} columns)"]))
        rows.append(("header", header))
        rows += [("data", r) for r in body]
        rows.append(("blank", []))
    return rows


def contents_rows(found: list[dict]) -> list[tuple[str, list]]:
    rows: list[tuple[str, list]] = [
        ("title", ["Supplementary and appendix tables"]),
        ("caption", [INTRO.replace("`", "")]),
        ("blank", []),
        ("header", ["Sheet", "Table", "What it holds", "Files"]),
    ]
    for group, heading in (("S", "Supplementary tables"), ("A", "Appendix tables")):
        rows.append(("section", [heading]))
        for t in (t for t in found if t["id"].startswith(group)):
            sizes = []
            for name in t["files"]:
                header, body = read_file(name)
                sizes.append(f"{name} ({len(body):,} rows, {len(header)} columns)")
            rows.append(("index", [t["id"], t["name"], t["caption"], "\n".join(sizes)]))
    return rows


def workbook_rows() -> dict[str, list[tuple[str, list]]]:
    found = tables()
    sheets = {"Contents": contents_rows(found)}
    for t in found:
        sheets[t["id"]] = table_rows(t)
    return sheets


# --------------------------------------------------------------------------------- writing


FONT = Font(name="Calibri", size=11)
BOLD = Font(name="Calibri", size=11, bold=True)
TITLE = Font(name="Calibri", size=14, bold=True)
LINK = Font(name="Calibri", size=11, color="0563C1", underline="single")
GREY = PatternFill("solid", fgColor="E7E6E6")
BLUE = PatternFill("solid", fgColor="DDEBF7")
BELOW = Border(bottom=Side(style="thin", color="7F7F7F"))
WRAP = Alignment(wrap_text=True, vertical="top")


def lines_needed(text: str, width: float) -> int:
    """Lines that a wrapped cell needs; the row height is set from it because a spreadsheet does not fit a merged cell."""
    return sum(max(1, math.ceil(len(line) / max(width * CHARS_PER_UNIT, 1))) for line in str(text).split("\n"))


def put(ws, r: int, c: int, value, font=FONT) -> None:
    cell = ws.cell(row=r, column=c, value=value)
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"  # a value that begins with "=" is text here, never a formula
    cell.font = font


def write_sheet(ws, rows: list[tuple[str, list]], widths: list[float]) -> None:
    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    span = min(CAPTION_COLUMNS, len(widths))
    for r, (kind, values) in enumerate(rows, start=1):
        for c, v in enumerate(values, start=1):
            put(ws, r, c, v, TITLE if kind == "title" else BOLD if kind in ("file", "header", "section") else FONT)
        if kind == "caption":
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=span)
            ws.cell(row=r, column=1).alignment = WRAP
            ws.row_dimensions[r].height = min(409, 15.5 * lines_needed(values[0], sum(widths[:span])) + 8)
        elif kind == "file":
            for c in range(1, len(widths) + 1):
                ws.cell(row=r, column=c).fill = BLUE
        elif kind == "header":
            for c in range(1, len(values) + 1):
                ws.cell(row=r, column=c).fill = GREY
                ws.cell(row=r, column=c).border = BELOW
        elif kind == "section":
            for c in range(1, len(widths) + 1):
                ws.cell(row=r, column=c).fill = BLUE
        elif kind == "index":
            for c in range(1, len(values) + 1):
                ws.cell(row=r, column=c).alignment = WRAP
            ws.row_dimensions[r].height = min(409, 15.5 * max(lines_needed(values[2], widths[2]), lines_needed(values[3], widths[3])) + 8)
            link = ws.cell(row=r, column=1)
            link.hyperlink = Hyperlink(ref=link.coordinate, location=f"'{values[0]}'!A1", display=values[0])
            link.font = LINK


def widths_for(sheet: str, rows: list[tuple[str, list]]) -> list[float]:
    if sheet == "Contents":
        return [9, 26, 100, 78]
    ncols = max(len(v) for k, v in rows if k in ("header", "data"))
    widest = [8.0] * ncols
    for kind, values in rows:
        if kind in ("header", "data"):
            for c, v in enumerate(values):
                widest[c] = max(widest[c], min(MAX_WIDTH, len(str(v)) + 2))
    return widest


def build() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.title = "Supplementary and appendix tables: the fastVEP ACMG/AMP classifier"
    wb.properties.creator = "fastVEP"
    for sheet, rows in workbook_rows().items():
        write_sheet(wb.create_sheet(sheet), rows, widths_for(sheet, rows))
    return wb


# --------------------------------------------------------------------------------- comparing


def norm(v):
    """A value as a spreadsheet keeps it: a number is a number whatever its Python type, and a blank is None."""
    if v is None:
        return None
    if isinstance(v, bool):
        return ("b", v)
    if isinstance(v, (int, float)):
        return ("n", float(v))
    return ("s", str(v))


def same(a, b) -> bool:
    if a is None or b is None or a[0] != b[0]:
        return a == b
    if a[0] == "n":
        return math.isclose(a[1], b[1], rel_tol=1e-15, abs_tol=0.0)
    return a == b


def grid(rows: list[list]) -> list[list]:
    out = [[norm(v) for v in r] for r in rows]
    for r in out:
        while r and r[-1] is None:
            r.pop()
    while out and not out[-1]:
        out.pop()
    return out


def differences() -> list[str]:
    """What differs between the workbook on disk and what data/ and manuscript.md now say; empty if nothing."""
    if not WORKBOOK.exists():
        return [f"{WORKBOOK.name} does not exist"]
    want = {s: grid([values for _, values in rows]) for s, rows in workbook_rows().items()}
    have_wb = load_workbook(WORKBOOK)
    have = {ws.title: grid([list(r) for r in ws.iter_rows(values_only=True)]) for ws in have_wb.worksheets}
    out = []
    if list(have) != list(want):
        out.append(f"sheets are {list(have)}, expected {list(want)}")
    for sheet in want:
        if sheet not in have:
            continue
        for i in range(max(len(have[sheet]), len(want[sheet]))):
            a = have[sheet][i] if i < len(have[sheet]) else []
            b = want[sheet][i] if i < len(want[sheet]) else []
            if len(a) != len(b) or not all(same(x, y) for x, y in zip(a, b)):
                out.append(f"sheet {sheet}, row {i + 1}: workbook has {a}, data and manuscript say {b}")
                break
    return out


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    mode = ap.add_mutually_exclusive_group(required=True)
    mode.add_argument("--write", action="store_true", help="write the workbook unless it is already up to date")
    mode.add_argument("--check", action="store_true", help="exit non-zero if the workbook is not up to date")
    args = ap.parse_args()
    diff = differences()
    if args.check:
        print(f"{WORKBOOK.name}: " + ("up to date" if not diff else f"{len(diff)} difference(s)"))
        for d in diff:
            print("DIFF", d)
        sys.exit(1 if diff else 0)
    if not diff:
        print(f"{WORKBOOK.name}: already up to date")
        return
    build().save(WORKBOOK)
    again = differences()
    if again:
        sys.exit("the workbook written does not read back as what was meant:\n" + "\n".join(again))
    print(f"wrote {WORKBOOK.name}: {len(workbook_rows())} sheets")


if __name__ == "__main__":
    main()
